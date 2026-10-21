"""
Build the Revision-4 point-by-point response letter (.docx).

Uses the Revision-3 letter (manuscript/source/Response_to_Reviewers_R3.docx) only as
a style template (page setup, fonts, the 'List Bullet' style); its text is replaced.
Every number is read from the same result files, through the same functions, as
manuscript/build_manuscript_r4.py, so the letter and the manuscript cannot disagree.

Usage (from the project root):  python manuscript/build_response_letter_r4.py
Output: manuscript/Response_to_Reviewers_R4.docx
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import build_manuscript_r4 as M  # noqa: E402
from docx_edit import parse_md  # noqa: E402

TEMPLATE = HERE / "source" / "Response_to_Reviewers_R3.docx"


def clear_body(doc):
    body = doc.element.body
    for el in list(body):
        if el.tag.endswith("}sectPr"):
            continue
        body.remove(el)


def para(doc, md, style=None):
    p = doc.add_paragraph(style=style)
    for text, bold, italic in parse_md(md):
        r = p.add_run(text)
        if bold:
            r.bold = True
        if italic:
            r.italic = True
    return p


def build(out_path, sections):
    doc = Document(str(TEMPLATE))
    clear_body(doc)
    for kind, md in sections:
        if kind == "p":
            para(doc, md)
        elif kind == "b":
            para(doc, md, "List Bullet")
        elif kind == "blank":
            doc.add_paragraph()
    doc.save(str(out_path))


def main():
    n = M.load()
    sections = M.response_sections(n)
    out = HERE / "Response_to_Reviewers_R4.docx"
    build(out, sections)
    print(f"Saved {out.relative_to(HERE.parent)}")


if __name__ == "__main__":
    main()
